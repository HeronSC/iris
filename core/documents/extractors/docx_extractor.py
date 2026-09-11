# File: core/documents/extractors/docx_extractor.py

from __future__ import annotations

import xml.etree.ElementTree as ET
from pathlib import Path

from core.documents.models import ExtractedDocument

_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


class DocxExtractor:

    name = "docx"

    def __init__(self, *, max_table_rows: int = 500) -> None:
        self.max_table_rows = max_table_rows

    def supports(self, path: Path) -> bool:
        return path.suffix.lower() == ".docx"

    def extract(self, path: Path) -> ExtractedDocument:
        try:
            #! @allow-local-import
            import docx
            #! @allow-local-import
            from docx.opc.exceptions import PackageNotFoundError
        except Exception:
            return ExtractedDocument(text="", content_status="text_unavailable", extractor=self.name, error="Install python-docx for Word extraction")

        try:
            document = docx.Document(str(path))
        except PackageNotFoundError as error:
            return ExtractedDocument(text="", content_status="error", extractor=self.name, error=f"Not a Word document: {error}")
        except Exception as error:
            return ExtractedDocument(text="", content_status="error", extractor=self.name, error=str(error))

        try:
            lines: list[str] = []
            headers = self._section_text(document, "header")
            if headers:
                lines.append("Header: " + " | ".join(headers))
            lines.extend(self._body_lines(document))
            footers = self._section_text(document, "footer")
            if footers:
                lines.append("Footer: " + " | ".join(footers))
            footnotes = self._footnotes(document)
            if footnotes:
                lines.append("Footnotes:")
                lines.extend(f"- {note}" for note in footnotes)
            if not lines:
                return ExtractedDocument(text="", content_status="text_unavailable", extractor=self.name, error="No text in document")
            return ExtractedDocument(text="\n".join(lines), content_status="indexed", extractor=self.name)
        except Exception as error:
            return ExtractedDocument(text="", content_status="error", extractor=self.name, error=str(error))

    def _body_lines(self, document: object) -> list[str]:
        lines: list[str] = []
        body = document.element.body
        #! @allow-local-import
        from docx.table import Table
        #! @allow-local-import
        from docx.text.paragraph import Paragraph

        for child in body.iterchildren():
            tag = child.tag
            if tag == f"{_W}p":
                text = Paragraph(child, document).text.strip()
                if text:
                    lines.append(text)
            elif tag == f"{_W}tbl":
                lines.extend(self._table_lines(Table(child, document)))
        return lines

    def _table_lines(self, table: object) -> list[str]:
        lines: list[str] = []
        for row_index, row in enumerate(table.rows):
            if row_index >= self.max_table_rows:
                lines.append(f"... table truncated after {self.max_table_rows} rows")
                break
            cells: list[str] = []
            previous = None
            for cell in row.cells:
                if cell._tc is previous:
                    continue
                previous = cell._tc
                cells.append(" ".join(cell.text.split()))
            if any(cells):
                lines.append(" | ".join(cells))
        return lines

    @staticmethod
    def _section_text(document: object, part: str) -> list[str]:
        seen: list[str] = []
        for section in document.sections:
            container = getattr(section, part, None)
            if container is None or getattr(container, "is_linked_to_previous", False):
                continue
            for paragraph in container.paragraphs:
                text = " ".join(paragraph.text.split())
                if text and text not in seen:
                    seen.append(text)
            for table in getattr(container, "tables", []):
                for row in table.rows:
                    text = " | ".join(" ".join(cell.text.split()) for cell in row.cells)
                    if text.strip(" |") and text not in seen:
                        seen.append(text)
        return seen

    @staticmethod
    def _footnotes(document: object) -> list[str]:
        notes: list[str] = []
        try:
            for relationship in document.part.rels.values():
                if not relationship.reltype.endswith("/footnotes"):
                    continue
                root = ET.fromstring(relationship.target_part.blob)
                for note in root.iter(f"{_W}footnote"):
                    if note.get(f"{_W}type") in {"separator", "continuationSeparator"}:
                        continue
                    text = " ".join(t.text for t in note.iter(f"{_W}t") if t.text)
                    text = " ".join(text.split())
                    if text:
                        notes.append(text)
        except Exception:
            return notes
        return notes
