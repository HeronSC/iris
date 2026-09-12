# File: core/documents/extractors/excel_extractor.py

from __future__ import annotations

from datetime import date, datetime, time
from pathlib import Path
from typing import Any

from core.documents.models import ExtractedDocument


class ExcelExtractor:

    name = "excel"

    def __init__(self, *, max_rows_per_sheet: int = 2000, max_cells: int = 40000) -> None:
        self.max_rows_per_sheet = max_rows_per_sheet
        self.max_cells = max_cells

    def supports(self, path: Path) -> bool:
        return path.suffix.lower() in {".xlsx", ".xlsm", ".xls"}

    def extract(self, path: Path) -> ExtractedDocument:
        if path.suffix.lower() == ".xls":
            return ExtractedDocument(text="", content_status="text_unavailable", extractor=self.name, error="Legacy .xls is not supported; save as .xlsx")
        try:
            #! @allow-local-import
            from openpyxl import load_workbook
        except ImportError:
            return ExtractedDocument(text="", content_status="text_unavailable", extractor=self.name, error="Install openpyxl for Excel extraction")

        try:
            workbook = load_workbook(str(path), read_only=True, data_only=True)
        except (OSError, ValueError, RuntimeError, TypeError) as error:
            return ExtractedDocument(text="", content_status="error", extractor=self.name, error=str(error))

        try:
            lines: list[str] = []
            names = list(workbook.sheetnames)
            if names:
                lines.append("Sheets: " + ", ".join(names))
            cells_seen = 0
            truncated = False
            for sheet in workbook.worksheets:
                lines.append(f"Worksheet: {sheet.title}")
                for row_index, row in enumerate(sheet.iter_rows(values_only=True)):
                    if row_index >= self.max_rows_per_sheet:
                        lines.append(f"... sheet truncated after {self.max_rows_per_sheet} rows")
                        truncated = True
                        break
                    rendered = [self._render(value) for value in row]
                    while rendered and rendered[-1] == "":
                        rendered.pop()
                    if not any(rendered):
                        continue
                    cells_seen += len(rendered)
                    lines.append(" | ".join(rendered))
                    if cells_seen >= self.max_cells:
                        lines.append(f"... workbook truncated after {self.max_cells} cells")
                        truncated = True
                        break
                if truncated:
                    break
            if len(lines) <= 1:
                return ExtractedDocument(text="\n".join(lines), content_status="text_unavailable", extractor=self.name, error="Workbook has no cell values")
            return ExtractedDocument(
                text="\n".join(lines),
                content_status="indexed",
                extractor=self.name,
                error="Workbook was truncated" if truncated else None,
            )
        except (OSError, ValueError, RuntimeError, TypeError) as error:
            return ExtractedDocument(text="", content_status="error", extractor=self.name, error=str(error))
        finally:
            try:
                workbook.close()
            except (OSError, ValueError, RuntimeError, TypeError):
                pass

    @staticmethod
    def _render(value: Any) -> str:
        if value is None:
            return ""
        if isinstance(value, datetime):
            return value.date().isoformat() if value.time() == time(0, 0) else value.isoformat(sep=" ", timespec="minutes")
        if isinstance(value, date):
            return value.isoformat()
        if isinstance(value, bool):
            return "TRUE" if value else "FALSE"
        if isinstance(value, float) and value.is_integer():
            return str(int(value))
        return " ".join(str(value).split())
