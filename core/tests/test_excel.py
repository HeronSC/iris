# File: core/tests/test_excel.py

"""Excel (4.1), read-only: the live workbook through COM, closed files through openpyxl, one service over both."""

from __future__ import annotations

import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace

from openpyxl import Workbook
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.table import Table

from core.actions.implementations.excel_tools import ExcelCheckAction, ExcelFindAction, ExcelReadAction, ExcelWorkbookAction
from core.actions.models import ActionRequest
from core.excel import ExcelService, WorkbookUnavailable
from core.excel.closed import ClosedExcel
from core.excel.live import LiveExcel
from core.excel.models import cell_address, error_text
from core.results.models import ResultKind

BUDGET = "budget.xlsx"


def _build_workbook(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Q3"
    sheet.append(["Item", "Qty", "Price", "Total"])
    sheet.append(["Widget", 2, 5.5, "=B2*C2"])
    sheet.append(["Gadget", 0, 3, "=C3/B3"])
    sheet.append(["Sum", None, None, "=SUM(D2:D3)"])
    sheet.add_table(Table(displayName="Sales", ref="A1:D4"))
    notes = workbook.create_sheet("Notes")
    notes["A1"] = "Review the widget price"
    workbook.defined_names["TotalCell"] = DefinedName("TotalCell", attr_text="Q3!$D$4")
    workbook.save(str(path))


class _FakeRange:
    def __init__(self, values, formulas=None, address="A1:D3", row=1, column=1):
        self.Value2 = values
        self.Formula = formulas if formulas is not None else values
        self._address = address
        self.Row = row
        self.Column = column
        self.Rows = SimpleNamespace(Count=len(values) if isinstance(values, (list, tuple)) else 1)
        self.Columns = SimpleNamespace(Count=len(values[0]) if isinstance(values, (list, tuple)) and values and isinstance(values[0], (list, tuple)) else 1)

    def Address(self, row_absolute, column_absolute):
        return self._address


class _FakeSheet:
    def __init__(self, name, values, formulas=None, tables=(), pivots=0):
        self.Name = name
        self.UsedRange = _FakeRange(values, formulas)
        self.ListObjects = [SimpleNamespace(Name=table) for table in tables]
        self._pivots = pivots
        self.Visible = -1
        self.ProtectContents = False

    def PivotTables(self):
        return SimpleNamespace(Count=self._pivots)

    def Range(self, address):
        return _FakeRange(((self.UsedRange.Value2[0][0],),), address=address)


class _FakeWorkbook:
    def __init__(self, full_name, sheets, saved=True):
        self.FullName = full_name
        self.Name = Path(full_name).name
        self._sheets = sheets
        self.ActiveSheet = sheets[0]
        self.Saved = saved
        self.ReadOnly = False
        self.Names = [SimpleNamespace(Name="TotalCell", RefersTo="=Q3!$D$4")]

    @property
    def Worksheets(self):
        class _Collection(list):
            def __call__(inner, name):
                for sheet in inner:
                    if sheet.Name == name:
                        return sheet
                raise KeyError(name)

        return _Collection(self._sheets)


class _FakeApp:
    def __init__(self, workbook):
        self.Workbooks = [workbook]
        self.ActiveWorkbook = workbook
        self.Selection = SimpleNamespace(Address=lambda r, c: "B2:B3")


def _live_fixture(full_name: str) -> _FakeApp:
    sheet = _FakeSheet("Q3", ((("Item", "Qty", "Total"), ("Widget", 2.0, 11.0), ("Gadget", 0.0, -2146826281))), formulas=((("Item", "Qty", "Total"), ("Widget", 2, "=B2*5.5"), ("Gadget", 0, "=C3/B3"))), tables=("Sales",), pivots=1)
    notes = _FakeSheet("Notes", "Review the widget price")
    return _FakeApp(_FakeWorkbook(full_name, [sheet, notes], saved=False))


class ModelTests(unittest.TestCase):
    def test_addresses_and_error_codes(self) -> None:
        self.assertEqual(cell_address(1, 1), "A1")
        self.assertEqual(cell_address(10, 28), "AB10")
        self.assertEqual(error_text(-2146826281), "#DIV/0!")
        self.assertEqual(error_text("#ref!"), "#REF!")
        self.assertIsNone(error_text(5))
        self.assertIsNone(error_text(True))


class ClosedTests(unittest.TestCase):
    def setUp(self) -> None:
        self.tempdir = tempfile.TemporaryDirectory()
        self.path = Path(self.tempdir.name) / BUDGET
        _build_workbook(self.path)
        self.closed = ClosedExcel()

    def tearDown(self) -> None:
        self.tempdir.cleanup()

    def test_describe_lists_sheets_tables_and_names(self) -> None:
        info = self.closed.describe(self.path)
        self.assertFalse(info.live)
        self.assertEqual([sheet.name for sheet in info.sheets], ["Q3", "Notes"])
        self.assertEqual((info.sheets[0].rows, info.sheets[0].columns), (4, 4))
        self.assertEqual(info.sheets[0].tables, ("Sales",))
        self.assertEqual(info.names, {"TotalCell": "Q3!$D$4"})
        self.assertIn("2 sheets: Q3 (4 x 4, tables Sales), Notes (1 x 1)", info.describe())
        self.assertIn("Named ranges: TotalCell = Q3!$D$4", info.describe())

    def test_read_gives_values_and_formulas_and_find_locates_text(self) -> None:
        data = self.closed.read(self.path, "Q3", "A1:D2", formulas=True)
        self.assertEqual(data.rows, (("Item", "Qty", "Price", "Total"), ("Widget", 2, 5.5, None)))
        self.assertEqual(data.formulas[1][3], "=B2*C2")
        whole = self.closed.read(self.path, None, None, max_rows=2)
        self.assertEqual(whole.sheet, "Q3")
        self.assertEqual(whole.address, "A1:D4")
        self.assertTrue(whole.truncated)
        hits = self.closed.find(self.path, "widget")
        self.assertEqual([(hit.sheet, hit.address) for hit in hits], [("Q3", "A2"), ("Notes", "A1")])
        self.assertEqual(self.closed.find(self.path, "widget", "Notes")[0].address, "A1")

    def test_errors_need_cached_values_so_a_fresh_openpyxl_file_reports_none(self) -> None:
        self.assertEqual(self.closed.errors(self.path), [])

    def test_missing_and_old_xls_files_are_explained(self) -> None:
        with self.assertRaises(WorkbookUnavailable):
            self.closed.describe(Path(self.tempdir.name) / "missing.xlsx")
        with self.assertRaises(WorkbookUnavailable) as caught:
            self.closed.describe(Path(self.tempdir.name) / "old.xls")
        self.assertIn(".xls", str(caught.exception))


class LiveTests(unittest.TestCase):
    def setUp(self) -> None:
        self.app = _live_fixture("C:\\Books\\budget.xlsx")
        self.live = LiveExcel(lambda: self.app)

    def test_describe_reads_sheets_names_selection_and_unsaved_state(self) -> None:
        info = self.live.describe(self.live.workbook(None))
        self.assertTrue(info.live)
        self.assertFalse(info.saved)
        self.assertEqual(info.selection, "B2:B3")
        self.assertEqual(info.active_sheet, "Q3")
        self.assertEqual(info.sheets[0].tables, ("Sales",))
        self.assertEqual(info.sheets[0].pivots, 1)
        self.assertEqual(info.names, {"TotalCell": "Q3!$D$4"})
        self.assertIn("Has unsaved changes.", info.describe())
        self.assertEqual(self.live.workbooks(), [("budget.xlsx", "C:\\Books\\budget.xlsx")])
        self.assertIsNotNone(self.live.workbook("BUDGET.XLSX"))
        self.assertIsNone(self.live.workbook("other.xlsx"))
        self.assertIsNone(LiveExcel(lambda: None).workbook(None))

    def test_read_find_and_errors_map_com_values(self) -> None:
        workbook = self.live.workbook(None)
        data = self.live.read(workbook, "Q3", None, formulas=True)
        self.assertEqual(data.rows[1], ("Widget", 2, 11))
        self.assertEqual(data.rows[2][2], "#DIV/0!")
        self.assertEqual(data.formulas[2][2], "=C3/B3")
        hits = self.live.find(workbook, "gadget")
        self.assertEqual([(hit.sheet, hit.address, hit.value) for hit in hits], [("Q3", "A3", "Gadget")])
        errors = self.live.errors(workbook)
        self.assertEqual([(hit.sheet, hit.address, hit.value, hit.formula) for hit in errors], [("Q3", "C3", "#DIV/0!", "=C3/B3")])


class ServiceAndToolTests(unittest.TestCase):
    def setUp(self) -> None:
        self.tempdir = tempfile.TemporaryDirectory()
        self.root = Path(self.tempdir.name)
        self.path = self.root / BUDGET
        _build_workbook(self.path)
        self.app = _live_fixture(str(self.root / "live.xlsx"))
        self.context = SimpleNamespace(paused=False, current=lambda: SimpleNamespace(target=str(self.root / "live.xlsx"), target_kind="workbook"))
        self.service = ExcelService([self.root], live=LiveExcel(lambda: self.app), context_service=self.context)
        self.execution = SimpleNamespace(excel_service=self.service)

    def tearDown(self) -> None:
        self.tempdir.cleanup()

    def test_resolution_prefers_the_workbook_in_view_then_open_names_then_disk(self) -> None:
        in_view = self.service.resolve(None)
        self.assertTrue(in_view.live)
        self.assertEqual(in_view.path, str(self.root / "live.xlsx"))
        self.assertTrue(self.service.resolve("live.xlsx").live)
        on_disk = self.service.resolve(str(self.path))
        self.assertFalse(on_disk.live)
        with self.assertRaises(WorkbookUnavailable):
            self.service.resolve(str(Path(self.tempdir.name).parent / "elsewhere.xlsx"))
        with self.assertRaises(WorkbookUnavailable):
            self.service.resolve("notopen.xlsx")
        nothing = ExcelService([self.root], live=LiveExcel(lambda: None))
        with self.assertRaises(WorkbookUnavailable) as caught:
            nothing.resolve(None)
        self.assertIn("No workbook is in view", str(caught.exception))

    def test_workbook_tool_describes_live_and_closed(self) -> None:
        action = ExcelWorkbookAction()
        validation = action.validate(ActionRequest(action="excel_workbook", arguments={}), self.execution)
        self.assertTrue(validation.ok)
        live = action.execute(ActionRequest(action="excel_workbook", arguments=validation.resolved_arguments), self.execution)
        self.assertEqual(live.status, "success")
        self.assertIn("open in Excel", live.message)
        self.assertEqual([item.kind for item in live.results], [ResultKind.TEXT, ResultKind.TABLE, ResultKind.TABLE, ResultKind.STATUS])
        closed = action.execute(ActionRequest(action="excel_workbook", arguments={"path": str(self.path)}), self.execution)
        self.assertIn("read from disk", closed.message)
        self.assertIn("tables Sales", closed.message)

    def test_read_find_and_check_tools(self) -> None:
        read = ExcelReadAction()
        validation = read.validate(ActionRequest(action="excel_read", arguments={"path": str(self.path), "sheet": "Q3", "range": "A1:B2", "formulas": True}), self.execution)
        self.assertTrue(validation.ok)
        result = read.execute(ActionRequest(action="excel_read", arguments=validation.resolved_arguments), self.execution)
        self.assertEqual(result.status, "success")
        self.assertIn("sheet Q3, range A1:B2: 2 rows", result.message)
        self.assertEqual(result.results[0].data["columns"], ["A", "B"])
        self.assertEqual(result.results[0].data["rows"][1], ["Widget", 2])
        self.assertEqual(result.results[1].title, "Formulas in Q3!A1:B2")
        find_action = ExcelFindAction()
        find_validation = find_action.validate(ActionRequest(action="excel_find", arguments={"text": "widget", "path": str(self.path)}), self.execution)
        self.assertTrue(find_validation.ok)
        find = find_action.execute(ActionRequest(action="excel_find", arguments=find_validation.resolved_arguments), self.execution)
        self.assertIn("2 cells in budget.xlsx contain 'widget'", find.message)
        check_live = ExcelCheckAction().execute(ActionRequest(action="excel_check", arguments={"path": "live.xlsx"}), self.execution)
        self.assertIn("Q3!C3: #DIV/0! from =C3/B3", check_live.message)
        self.assertEqual(check_live.results[0].kind, ResultKind.STATUS)
        check_closed = ExcelCheckAction().execute(ActionRequest(action="excel_check", arguments={"path": str(self.path)}), self.execution)
        self.assertIn("No error values", check_closed.message)
        missing = read.validate(ActionRequest(action="excel_read", arguments={"path": "nowhere.xlsx"}), self.execution)
        self.assertFalse(missing.ok)
        for action in (ExcelWorkbookAction(), ExcelReadAction(), ExcelFindAction(), ExcelCheckAction()):
            self.assertEqual(action.definition.permission.value, "read")
        no_service = ExcelWorkbookAction().validate(ActionRequest(action="excel_workbook", arguments={}), SimpleNamespace())
        self.assertFalse(no_service.ok)


if __name__ == "__main__":
    unittest.main()
