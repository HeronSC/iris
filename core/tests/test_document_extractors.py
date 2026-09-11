# File: core/tests/test_document_extractors.py

from __future__ import annotations

import base64
import datetime as dt
import tempfile
import unittest
from pathlib import Path

from core.documents.extractors.docx_extractor import DocxExtractor
from core.documents.extractors.excel_extractor import ExcelExtractor
from core.documents.extractors.pdf_extractor import PdfExtractor
from core.documents.ocr import OcrService, VisionOcr
from core.llm.models import LLMRequest, LLMResponse


class FakeOcr:
    name = "fake-ocr"

    def __init__(self, text: str = "SCANNED LINE", fail: bool = False) -> None:
        self.text = text
        self.fail = fail
        self.images: list[bytes] = []

    def read(self, image_png: bytes) -> str:
        self.images.append(image_png)
        if self.fail:
            raise RuntimeError("engine missing")
        return self.text


class PdfExtractorTests(unittest.TestCase):
    def setUp(self) -> None:
        self._tmp = tempfile.TemporaryDirectory()
        self.path = Path(self._tmp.name) / "sample.pdf"
        import pymupdf

        document = pymupdf.open()
        page = document.new_page()
        page.insert_text((72, 72), "Quarterly revenue grew 12 percent.")
        page = document.new_page()
        page.insert_text((72, 72), "Costs were flat.")
        document.new_page()  # a blank page: a scan with no text layer
        document.save(str(self.path))
        document.close()

    def tearDown(self) -> None:
        self._tmp.cleanup()

    def test_pages_are_marked_for_citation(self) -> None:
        result = PdfExtractor().extract(self.path)
        self.assertEqual(result.content_status, "indexed")
        self.assertIn("[page 1]\nQuarterly revenue grew 12 percent.", result.text)
        self.assertIn("[page 2]\nCosts were flat.", result.text)
        self.assertNotIn("[page 3]", result.text)

    def test_blank_pages_go_to_ocr_in_page_order(self) -> None:
        ocr = FakeOcr("Handwritten note")
        result = PdfExtractor(ocr=ocr).extract(self.path)
        self.assertEqual(len(ocr.images), 1)
        self.assertTrue(ocr.images[0].startswith(b"\x89PNG"))
        self.assertIn("[page 3] (ocr)\nHandwritten note", result.text)
        self.assertLess(result.text.index("[page 2]"), result.text.index("[page 3]"))

    def test_scan_only_document_without_ocr_says_why(self) -> None:
        import pymupdf

        scan = Path(self._tmp.name) / "scan.pdf"
        document = pymupdf.open()
        document.new_page()
        document.save(str(scan))
        document.close()
        result = PdfExtractor().extract(scan)
        self.assertEqual(result.content_status, "text_unavailable")
        self.assertIn("no OCR reader configured", result.error or "")

        failing = PdfExtractor(ocr=FakeOcr(fail=True)).extract(scan)
        self.assertEqual(failing.content_status, "text_unavailable")
        self.assertIn("OCR failed", failing.error or "")

    def test_not_a_pdf_is_an_error_not_a_crash(self) -> None:
        bogus = Path(self._tmp.name) / "bogus.pdf"
        bogus.write_bytes(b"not a pdf")
        result = PdfExtractor().extract(bogus)
        self.assertEqual(result.content_status, "error")


class DocxExtractorTests(unittest.TestCase):
    def setUp(self) -> None:
        self._tmp = tempfile.TemporaryDirectory()
        self.path = Path(self._tmp.name) / "sample.docx"
        import docx

        document = docx.Document()
        document.sections[0].header.paragraphs[0].text = "Acme Internal"
        document.sections[0].footer.paragraphs[0].text = "Page footer text"
        document.add_heading("Project Plan", level=1)
        document.add_paragraph("The first milestone is the prototype.")
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Task"
        table.cell(0, 1).text = "Owner"
        table.cell(1, 0).text = "Prototype"
        table.cell(1, 1).text = "Henry"
        document.add_paragraph("Closing remarks.")
        document.save(str(self.path))

    def tearDown(self) -> None:
        self._tmp.cleanup()

    def test_body_tables_headers_and_footers_in_order(self) -> None:
        result = DocxExtractor().extract(self.path)
        self.assertEqual(result.content_status, "indexed", result.error)
        lines = result.text.splitlines()
        self.assertEqual(lines[0], "Header: Acme Internal")
        self.assertIn("Project Plan", lines)
        self.assertIn("Task | Owner", lines)
        self.assertIn("Prototype | Henry", lines)
        self.assertLess(lines.index("The first milestone is the prototype."), lines.index("Task | Owner"))
        self.assertLess(lines.index("Prototype | Henry"), lines.index("Closing remarks."))
        self.assertEqual(lines[-1], "Footer: Page footer text")

    def test_garbage_is_an_error(self) -> None:
        bogus = Path(self._tmp.name) / "bogus.docx"
        bogus.write_bytes(b"nope")
        self.assertEqual(DocxExtractor().extract(bogus).content_status, "error")


class ExcelExtractorTests(unittest.TestCase):
    def setUp(self) -> None:
        self._tmp = tempfile.TemporaryDirectory()
        self.path = Path(self._tmp.name) / "sample.xlsx"
        from openpyxl import Workbook

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Sales"
        sheet.append(["Date", "Region", "Amount", "Closed"])
        sheet.append([dt.datetime(2026, 9, 1), "East", 1250.0, True])
        sheet.append([dt.date(2026, 9, 2), "West", 99.5, False])
        sheet.append([None, None, None, None])
        sheet.merge_cells("A6:B6")
        sheet["A6"] = "Merged title"
        other = workbook.create_sheet("Notes")
        other["A1"] = "  spaced   text  "
        workbook.save(str(self.path))

    def tearDown(self) -> None:
        self._tmp.cleanup()

    def test_values_dates_and_sheets(self) -> None:
        result = ExcelExtractor().extract(self.path)
        self.assertEqual(result.content_status, "indexed", result.error)
        lines = result.text.splitlines()
        self.assertEqual(lines[0], "Sheets: Sales, Notes")
        self.assertIn("Worksheet: Sales", lines)
        self.assertIn("Date | Region | Amount | Closed", lines)
        self.assertIn("2026-09-01 | East | 1250 | TRUE", lines)
        self.assertIn("2026-09-02 | West | 99.5 | FALSE", lines)
        self.assertIn("Merged title", lines)
        self.assertIn("Worksheet: Notes", lines)
        self.assertIn("spaced text", lines)
        self.assertIsNone(result.error)

    def test_truncation_is_reported(self) -> None:
        result = ExcelExtractor(max_rows_per_sheet=2).extract(self.path)
        self.assertEqual(result.error, "Workbook was truncated")
        self.assertIn("... sheet truncated after 2 rows", result.text)

    def test_old_xls_is_declined_clearly(self) -> None:
        legacy = Path(self._tmp.name) / "old.xls"
        legacy.write_bytes(b"\xd0\xcf\x11\xe0")
        result = ExcelExtractor().extract(legacy)
        self.assertEqual(result.content_status, "text_unavailable")
        self.assertIn(".xlsx", result.error or "")


class FakeVisionClient:
    def __init__(self, text: str = "Seen by the model") -> None:
        self.text = text
        self.requests: list[LLMRequest] = []

    def chat(self, request: LLMRequest) -> LLMResponse:
        self.requests.append(request)
        return LLMResponse(content=self.text, model="qwen2.5vl:7b")


class OcrServiceTests(unittest.TestCase):
    def test_first_reader_with_text_wins(self) -> None:
        empty = FakeOcr("")
        good = FakeOcr("second reader text")
        service = OcrService([empty, good])
        self.assertEqual(service.read(b"png"), "second reader text")
        self.assertEqual(service.last_reader, "fake-ocr")
        self.assertEqual(len(empty.images), 1)

    def test_a_failing_reader_is_skipped_and_all_failing_raises(self) -> None:
        service = OcrService([FakeOcr(fail=True), FakeOcr("rescued")])
        self.assertEqual(service.read(b"png"), "rescued")
        with self.assertRaises(RuntimeError):
            OcrService([FakeOcr(fail=True)]).read(b"png")
        self.assertFalse(OcrService([None]).available)

    def test_vision_reader_sends_the_image_on_the_vision_task(self) -> None:
        client = FakeVisionClient()
        reader = VisionOcr(client)
        self.assertEqual(reader.read(b"\x89PNGdata"), "Seen by the model")
        request = client.requests[-1]
        self.assertEqual(request.task, "vision")
        self.assertEqual(request.think, False)
        message = request.messages[0]
        self.assertEqual(message.images, (b"\x89PNGdata",))
        self.assertEqual(message.to_ollama()["images"], [base64.b64encode(b"\x89PNGdata").decode("ascii")])
        self.assertIn("Transcribe", message.content)


if __name__ == "__main__":
    unittest.main()
