# File: core/excel/closed.py

from __future__ import annotations

import logging
from pathlib import Path
from typing import Any

from core.excel.models import MAX_SCAN_CELLS, CellHit, RangeData, SheetInfo, WorkbookInfo, cell_address, error_text

logger = logging.getLogger(__name__)


class WorkbookUnavailable(RuntimeError):
    pass


def _clean(value: Any) -> Any:
    if hasattr(value, "isoformat"):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return value


class ClosedExcel:
    def _open(self, path: Path, *, data_only: bool):
        try:
            #! @allow-local-import
            from openpyxl import load_workbook
        except ImportError as error:
            raise WorkbookUnavailable("openpyxl is not installed") from error
        if not path.is_file():
            raise WorkbookUnavailable(f"{path} does not exist")
        if path.suffix.lower() == ".xls":
            raise WorkbookUnavailable("Legacy .xls is not supported; save it as .xlsx")
        try:
            return load_workbook(str(path), data_only=data_only, read_only=False)
        except PermissionError as error:
            raise WorkbookUnavailable(f"{path.name} is locked by another program (Excel or OneDrive); close it or wait for sync") from error
        except Exception as error:
            raise WorkbookUnavailable(f"Could not open {path.name}: {error}") from error

    def describe(self, path: Path) -> WorkbookInfo:
        workbook = self._open(path, data_only=True)
        try:
            sheets: list[SheetInfo] = []
            for sheet in workbook.worksheets:
                has_cells = sheet.max_row > 1 or sheet.max_column > 1 or sheet.cell(1, 1).value is not None
                pivots = len(getattr(sheet, "_pivots", []) or [])
                sheets.append(
                    SheetInfo(
                        name=str(sheet.title),
                        rows=int(sheet.max_row) if has_cells else 0,
                        columns=int(sheet.max_column) if has_cells else 0,
                        tables=tuple(str(name) for name in sheet.tables.keys()),
                        pivots=pivots,
                        hidden=str(sheet.sheet_state) != "visible",
                        protected=bool(getattr(getattr(sheet, "protection", None), "sheet", False)),
                    )
                )
            names: dict[str, str] = {}
            defined = getattr(workbook, "defined_names", {})
            items = defined.items() if hasattr(defined, "items") else [(item.name, item) for item in defined.definedName]
            for name, definition in items:
                names[str(name)] = str(getattr(definition, "attr_text", definition))
            return WorkbookInfo(path=str(path), name=path.name, live=False, sheets=tuple(sheets), names=names, active_sheet=str(workbook.active.title) if workbook.active is not None else None)
        finally:
            workbook.close()

    def read(self, path: Path, sheet: str | None, address: str | None, *, formulas: bool = False, max_rows: int = 100) -> RangeData:
        values_book = self._open(path, data_only=True)
        formula_book = self._open(path, data_only=False) if formulas else None
        try:
            target = values_book[sheet] if sheet else values_book.active
            cells = target[address] if address else target.iter_rows(min_row=1, max_row=target.max_row, max_col=target.max_column)
            rows: list[tuple[Any, ...]] = []
            grid = list(cells) if address else list(cells)
            if grid and not isinstance(grid[0], tuple):
                grid = [tuple(grid)] if not hasattr(grid[0], "__iter__") or hasattr(grid[0], "value") else grid
            first_row = first_column = last_row = last_column = None
            for row in grid:
                row_cells = row if isinstance(row, tuple) else (row,)
                rows.append(tuple(_clean(cell.value) for cell in row_cells))
                if row_cells:
                    first_row = row_cells[0].row if first_row is None else first_row
                    first_column = row_cells[0].column if first_column is None else first_column
                    last_row = row_cells[-1].row
                    last_column = row_cells[-1].column
            truncated = len(rows) > max_rows
            formula_rows: tuple[tuple[str, ...], ...] = ()
            if formula_book is not None:
                source = formula_book[str(target.title)]
                formula_cells = source[address] if address else source.iter_rows(min_row=1, max_row=source.max_row, max_col=source.max_column)
                collected: list[tuple[str, ...]] = []
                for row in formula_cells:
                    row_cells = row if isinstance(row, tuple) else (row,)
                    collected.append(tuple("" if cell.value is None else str(cell.value) for cell in row_cells))
                formula_rows = tuple(collected[:max_rows])
            resolved = address or (f"{cell_address(first_row, first_column)}:{cell_address(last_row, last_column)}" if first_row and last_row else "A1")
            return RangeData(sheet=str(target.title), address=resolved, rows=tuple(rows[:max_rows]), formulas=formula_rows, truncated=truncated)
        finally:
            values_book.close()
            if formula_book is not None:
                formula_book.close()

    def find(self, path: Path, text: str, sheet: str | None = None, *, limit: int = 50) -> list[CellHit]:
        needle = text.strip().casefold()
        workbook = self._open(path, data_only=True)
        try:
            hits: list[CellHit] = []
            scanned = 0
            targets = [workbook[sheet]] if sheet else list(workbook.worksheets)
            for target in targets:
                for row in target.iter_rows():
                    for cell in row:
                        scanned += 1
                        if scanned > MAX_SCAN_CELLS:
                            return hits
                        if cell.value is None:
                            continue
                        if needle in str(_clean(cell.value)).casefold():
                            hits.append(CellHit(sheet=str(target.title), address=cell.coordinate, value=_clean(cell.value)))
                            if len(hits) >= limit:
                                return hits
            return hits
        finally:
            workbook.close()

    def errors(self, path: Path, *, limit: int = 100) -> list[CellHit]:
        values_book = self._open(path, data_only=True)
        formula_book = self._open(path, data_only=False)
        try:
            hits: list[CellHit] = []
            scanned = 0
            for target in values_book.worksheets:
                source = formula_book[str(target.title)]
                for row in target.iter_rows():
                    for cell in row:
                        scanned += 1
                        if scanned > MAX_SCAN_CELLS:
                            return hits
                        marker = error_text(cell.value)
                        if marker is None:
                            continue
                        formula = source[cell.coordinate].value
                        hits.append(CellHit(sheet=str(target.title), address=cell.coordinate, value=marker, formula=str(formula) if formula is not None else None))
                        if len(hits) >= limit:
                            return hits
            return hits
        finally:
            values_book.close()
            formula_book.close()


__all__ = ["ClosedExcel", "WorkbookUnavailable"]
